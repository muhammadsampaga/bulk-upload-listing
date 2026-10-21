"""Run with: python test_mass_upload.py (no network or backend writes)."""
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

import app as module


# Independent positional fixture: rumah123-listing PR #7553, massUploadHeaderCols.
BASE = 'no|tipe properti|kategori|jenis properti|id area|'
TAIL = '|Status|aktivasi premier|aktivasi featured|aktivasi sundul|jadwal sundul|durasi sundul|hari sundul|jam sundul|tanggal akhir sundul'
EXPECTED = {
    'rumah': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar tidur|kamar mandi|jumlah lantai|carport|sertifikat|kondisi properti|kondisi perabotan|selling point|hadap|fasilitas rumah|fasilitas perumahan|lebar jalan',
    'apartemen': 'luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar tidur|kamar mandi|nomor lantai|sertifikat|kondisi perabotan|kondisi properti|fasilitas lingkungan apartemen|selling point|tahun dibangun|pemandangan|fasilitas apartemen|peralatan',
    'tanah': 'luas tanah|satuan tanah|harga|satuan harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|sertifikat|selling point|panjang tanah|lebar tanah',
    'ruko': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar mandi|jumlah lantai|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas ruko',
    'pabrik': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar mandi|jumlah lantai|carport|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas pabrik',
    'gudang': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar mandi|jumlah lantai|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas gudang',
    'perkantoran': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar mandi|jumlah lantai|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas perkantoran',
    'ruang usaha': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar mandi|jumlah lantai|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas ruang usaha',
    'villa': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar tidur|kamar mandi|jumlah lantai|carport|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas kamar villa',
    'kost': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar tidur|kamar mandi|jumlah lantai|carport|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas kamar kost|fasilitas kost',
    'hotel': 'luas tanah|luas bangunan|harga|periode sewa|judul iklan|deskripsi iklan|gambar cover utama|kamar tidur|kamar mandi|jumlah lantai|sertifikat|kondisi properti|kondisi perabotan|selling point|fasilitas sekitar hotel',
}
EXPECTED = {key: (BASE + value + TAIL).split('|') for key, value in EXPECTED.items()}
NEW = {'hadap', 'fasilitas rumah', 'fasilitas perumahan', 'lebar jalan', 'tahun dibangun',
       'pemandangan', 'fasilitas apartemen', 'peralatan', 'panjang tanah', 'lebar tanah',
       'fasilitas kamar villa', 'fasilitas kamar kost', 'fasilitas kost', 'fasilitas sekitar hotel',
       'aktivasi sundul', 'hari sundul', 'jam sundul', 'tanggal akhir sundul'}


def rejects(action, message):
    try:
        action()
    except ValueError as error:
        assert message in str(error), str(error)
    else:
        raise AssertionError('Expected rejection: ' + message)


def check_compatibility(folder):
    source = Path(module.EXCEL_FILE)
    with module.app.test_client() as client:
        assert client.get('/download/zip').status_code == 200
        response = client.post('/submit', data={'tipe_properti': 'ruang_usaha',
                               'id_area': '449', 'judul_iklan': 'Invalid schedule',
                               'aktivasi_sundul': 'Frequently', 'jadwal_sundul': '2 times daily'})
        assert response.status_code == 302
        assert not module.read_excel_data(), 'Invalid schedule must not be saved'
        response = client.post('/submit', data={'tipe_properti': 'ruang_usaha',
                               'id_area': '449', 'judul_iklan': 'New submission',
                               'gambar_cover': (BytesIO(b'test image'), 'cover.jpg'),
                               'jadwal_sundul': '1 time daily', 'durasi_sundul': '7d'})
        assert response.status_code == 302
        assert client.get('/load-data/ruang_usaha/1').get_json()['judul_iklan'] == 'New submission'
        assert client.get('/download/zip').status_code == 200
    # Next, replace only the temporary workbook with a genuine legacy-shaped fixture.
    wb = Workbook()
    wb.active.title = 'Data Properti'  # Legacy empty sheet must not reach the backend.
    for tipe, columns in EXPECTED.items():
        assert module.get_excel_columns(tipe) == columns, tipe
        headers = [col for col in columns if col not in NEW]
        if tipe == 'hotel':
            headers.insert(headers.index('sertifikat'), 'carport')
        ws = wb.create_sheet(tipe.replace(' ', '_'))
        ws.append(headers)
        values = {'no': 1, 'tipe properti': tipe.replace(' ', '_'), 'judul iklan': 'Original',
                  'Status': 'Tunda', 'aktivasi premier': 'Yes', 'aktivasi featured': 'No',
                  'jadwal sundul': '2 times daily', 'durasi sundul': '30 hari',
                  'gambar cover utama': 'cover.jpg', 'sertifikat': 'shm', 'carport': 2}
        ws.append([values.get(header, '') for header in headers])
        images = folder / tipe.replace(' ', '_') / '1'
        images.mkdir(parents=True, exist_ok=True)
        (images / 'cover.jpg').write_bytes(b'test image')
    wb.save(source)
    wb.close()
    before = source.read_bytes()
    exported = load_workbook(module.export_excel())
    assert set(exported.sheetnames) == set(EXPECTED)
    for tipe, headers in EXPECTED.items():
        ws = exported[tipe]
        assert [cell.value for cell in ws[1]] == headers
        row = dict(zip(headers, next(ws.iter_rows(min_row=2, values_only=True))))
        assert row['Status'] == 'Tunda' and row['aktivasi premier'] == 'Yes'
        assert row['aktivasi sundul'] == 'Frequently' and row['durasi sundul'] == '30d'
        assert row['sertifikat'] == 'shm' and row['gambar cover utama'] == 'cover.jpg'
    exported.close()
    assert source.read_bytes() == before, 'Export must not mutate legacy workbook'
    with patch.object(Workbook, 'save', side_effect=OSError('disk full')):
        try:
            module.update_excel_row('rumah', '1', {'judul_iklan': 'Must not overwrite'})
        except OSError:
            pass
        else:
            raise AssertionError('Expected save failure')
    assert source.read_bytes() == before, 'Failed save must not truncate original workbook'

    module.update_excel_row('rumah', '1', {'judul_iklan': 'Edited', 'hadap': 'Timur'})
    module.save_to_excel({'no': 2, 'tipe_properti': 'rumah', 'judul_iklan': 'New',
                         'gambar_cover_utama': 'cover.jpg'}, 'rumah')
    (folder / 'rumah' / '2').mkdir()
    (folder / 'rumah' / '2' / 'cover.jpg').write_bytes(b'test image')
    assert module.get_next_no('rumah') == 3
    with module.app.test_client() as client:
        loaded = client.get('/load-data/rumah/1').get_json()
        assert loaded['gambar_cover_utama'] == 'cover.jpg' and loaded['hadap'] == 'Timur'
        assert client.get('/load-data/ruang_usaha/1').status_code == 200
        response = client.get('/download/excel')
        assert response.status_code == 200
        downloaded = load_workbook(BytesIO(response.data))
        assert downloaded['rumah']['U2'].value == 'Timur'
        downloaded.close()
        response = client.post('/update-data', data={'tipe_properti': 'rumah', 'no': '1',
                               'id_area': '449', 'judul_iklan': 'Via form',
                               'jadwal_sundul': '1 time daily', 'durasi_sundul': '7d'})
        assert response.status_code == 302
        item = client.get('/load-data/rumah/1').get_json()
        assert item['gambar_cover_utama'] == 'cover.jpg' and item['hadap'] == 'Timur'
        assert item['aktivasi_sundul'] == 'Frequently'
        end = (datetime.now() + timedelta(days=7)).strftime('%Y-%m-%d')
        client.post('/update-data', data={'tipe_properti': 'ruang_usaha', 'no': '1',
                    'aktivasi_sundul': 'Day Selection', 'hari_sundul': 'Senin;Jumat',
                    'jam_sundul': '09:00', 'tanggal_akhir_sundul': end})
        item = client.get('/load-data/ruang_usaha/1').get_json()
        assert item['tanggal_akhir_sundul'] == end and item['hari_sundul'] == 'Senin;Jumat'
        assert client.get('/').status_code == 200

    archive = Path(module.create_zip())
    previous_zip = archive.read_bytes()
    with ZipFile(archive) as zipped:
        assert 'ruang usaha/1/cover.jpg' in zipped.namelist()
        assert not any('\\' in name for name in zipped.namelist())
    with patch.object(module, 'get_zip_filename', return_value=str(archive)):
        (folder / 'rumah' / '1' / 'cover.jpg').unlink()
        rejects(module.create_zip, 'cover tidak ditemukan')
        assert archive.read_bytes() == previous_zip, 'Failed ZIP must preserve previous archive'
    module.update_excel_row('rumah', '1', {'durasi_sundul': 'nonsense'})
    with module.app.test_client() as client:
        assert client.get('/download/zip').status_code == 302, 'Do not serve stale ZIP'
    rejects(module.export_excel, 'Durasi sundul')
    module.update_excel_row('rumah', '1', {'aktivasi_sundul': 'NONE'})
    wb = load_workbook(source)
    ws = wb['rumah']
    ws.cell(1, ws.max_column + 1, 'unknown field')
    ws.cell(2, ws.max_column, 'must not disappear')
    wb.save(source)
    wb.close()
    rejects(module.export_excel, 'unknown_field')
    wb = load_workbook(source)
    ws = wb['rumah']
    ws.delete_cols(ws.max_column)
    ws = wb.create_sheet('gedung')
    ws.append(['no', 'tipe properti'])
    ws.append([1, 'gedung'])
    wb.save(source)
    wb.close()
    rejects(module.export_excel, 'tidak didukung')


def check_sundul():
    assert module.normalize_sundul({})['aktivasi_sundul'] == 'NONE'
    assert module.normalize_sundul({'aktivasi_sundul': 'NONE', 'jadwal_sundul': 'invalid'})['jadwal_sundul'] == ''
    end = (datetime.now() + timedelta(days=7)).strftime('%Y-%m-%d')
    data = {'aktivasi_sundul': 'Day Selection', 'hari_sundul': 'senin;Jumat;Senin',
            'jam_sundul': '09:00', 'tanggal_akhir_sundul': end}
    normalized = module.normalize_sundul(data)
    assert normalized['hari_sundul'] == 'Senin;Jumat'
    assert normalized['tanggal_akhir_sundul'] == datetime.strptime(end, '%Y-%m-%d').strftime('%d-%m-%Y')
    for changed, message in [({'hari_sundul': 'Holiday'}, 'Hari sundul'),
                             ({'jam_sundul': '25:00'}, 'Jam sundul'),
                             ({'tanggal_akhir_sundul': '2000-01-01'}, 'sudah lewat'),
                             ({'tanggal_akhir_sundul': '2026-02-30'}, 'tidak valid')]:
        rejects(lambda: module.normalize_sundul(dict(data, **changed)), message)
    rejects(lambda: module.normalize_sundul({'jadwal_sundul': '1 time daily'}), 'Durasi sundul')
    rejects(lambda: module.normalize_sundul({'aktivasi_sundul': 'Yes'}), 'Aktivasi sundul')


if __name__ == '__main__':
    check_sundul()
    with TemporaryDirectory() as directory:
        folder = Path(directory)
        with patch.object(module, 'UPLOAD_FOLDER', directory), patch.object(module, 'EXCEL_FILE', str(folder / 'mass_upload_template.xlsx')):
            check_compatibility(folder)
    print('Mass-upload v3 checks passed (11 types, legacy data, edits, schedules, ZIP failures).')
